#!/usr/bin/env python3
"""
Normalize vendor price lists (xlsx / xlsm / PDF) into one import-ready CSV per vendor.

    python3 scripts/price-lists/normalize.py            # all vendors
    python3 scripts/price-lists/normalize.py bea nabco  # just these

Input:  price-lists/raw/<year>/<original vendor file>   (archived as received)
Output: price-lists/normalized/<distributor-id>.csv     (feed to POST /api/distributors/:id/import)
        price-lists/normalized/manifest.json            (counts, effective dates, warnings)

Every CSV has the same header, chosen so the server's fuzzy column mapper
(IMPORT_COLUMN_ALIASES in server.js) lands each column without a saved profile:

    Part Number, Description, UOM, List Price, Net Price, Category, Manufacturer Part, Notes, <vendor extras...>

"Net Price" is always OUR COST from that vendor. "List Price" is MSRP/list when the
vendor gives one. Extra vendor-specific columns (e.g. BEA's pre-tariff net, Record's
competing distributor prices) are kept as trailing columns; the importer captures
unmapped columns into item.extra so nothing is lost.

Requires: openpyxl (pip3 install openpyxl) and pdftotext (brew install poppler).
"""
import csv
import glob
import json
import os
import re
import subprocess
import sys
import warnings
from datetime import date

warnings.filterwarnings("ignore")  # openpyxl DrawingML noise
import openpyxl  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
RAW_DIR = os.path.join(ROOT, "price-lists", "raw")
OUT_DIR = os.path.join(ROOT, "price-lists", "normalized")

HEADER = ["Part Number", "Description", "UOM", "List Price", "Net Price", "Category", "Manufacturer Part", "Notes"]


# ───────────────────────── helpers ─────────────────────────

def find_raw(pattern):
    """Newest file under price-lists/raw/*/ whose name matches the glob (case-insensitive)."""
    hits = []
    for path in glob.glob(os.path.join(RAW_DIR, "*", "*")):
        if glob.fnmatch.fnmatch(os.path.basename(path).lower(), pattern.lower()):
            hits.append(path)
    if not hits:
        return None
    return max(hits, key=os.path.getmtime)


def money(v):
    """'$1,228.00 ' / 337 / '337.0' → 337.0 ; blank/CALL/'-' → None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip()
    if not s or s.upper() in ("CALL", "-", "N/A", "TBD"):
        return None
    s = re.sub(r"[^0-9.\-]", "", s)
    if s in ("", "-", ".", "-."):
        return None
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def part_key(s):
    """Part numbers arrive as floats (10.1311), ints (162460) or padded strings."""
    if isinstance(s, float) and s.is_integer():
        return str(int(s))
    return clean(s)


def sheet_rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def pdftotext(path, layout=True):
    args = ["pdftotext", "-layout" if layout else "-raw", path, "-"]
    return subprocess.run(args, capture_output=True, text=True, check=True).stdout


def ocr_pdf(path):
    """OCR with macOS Vision (scripts/price-lists/ocr-pdf.swift) → list of {page,x,y,w,h,text}, top-left origin."""
    script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ocr-pdf.swift")
    out = subprocess.run(["swift", script, path], capture_output=True, text=True, check=True).stdout
    return [json.loads(l) for l in out.splitlines() if l.strip()]


def write_csv(vendor_id, rows, extras=()):
    """rows: list of dicts keyed by HEADER names (+ extras). One row per part number: the first
    priced row wins; any later row with a *different* net price is recorded in Notes rather than lost."""
    os.makedirs(OUT_DIR, exist_ok=True)
    cols = HEADER + list(extras)
    by_key, order, dupes = {}, [], 0
    for r in rows:
        k = r["Part Number"].upper()
        if not k:
            continue
        if k not in by_key:
            by_key[k] = r
            order.append(k)
            continue
        dupes += 1
        keep = by_key[k]
        if keep.get("Net Price") is None and r.get("Net Price") is not None:
            r["Notes"] = clean(" ; ".join(x for x in (r.get("Notes"), keep.get("Notes")) if x))
            by_key[k] = r
        elif r.get("Net Price") is not None and r.get("Net Price") != keep.get("Net Price"):
            also = f"Also listed at ${r['Net Price']:.2f}" + (f" ({r['Description'][:40]})" if r.get("Description") != keep.get("Description") else "")
            if also not in (keep.get("Notes") or ""):
                keep["Notes"] = clean("; ".join(x for x in (keep.get("Notes"), also) if x))
    out = [by_key[k] for k in order]
    path = os.path.join(OUT_DIR, f"{vendor_id}.csv")
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in out:
            w.writerow({c: ("" if r.get(c) is None else r.get(c)) for c in cols})
    return path, len(out), dupes


def row(part, desc, uom="", lst=None, net=None, cat="", mfr="", notes="", **extra):
    d = {"Part Number": part_key(part), "Description": clean(desc), "UOM": clean(uom).upper(),
         "List Price": lst, "Net Price": net, "Category": clean(cat), "Manufacturer Part": clean(mfr),
         "Notes": clean(notes)}
    d.update(extra)
    return d


# ───────────────────────── vendors ─────────────────────────
# Each parser returns (rows, extras, meta). meta carries effective date / source notes.

def parse_bea():
    """BEA: Product | Description | Net Price | Actual Cost + 6% Tariff | MSRP.
    Our cost = tariff-inclusive column (what we actually pay); pre-tariff net kept as an extra."""
    path = find_raw("*BEA*PRICE LIST*.xlsx")
    rows, hdr = [], None
    for r in sheet_rows(path):
        cells = [clean(c) for c in r]
        if hdr is None:
            if "Product" in cells and "MSRP" in cells:
                hdr = {name: cells.index(name) for name in cells if name}
            continue
        part = r[hdr["Product"]]
        if part is None or not clean(part):
            continue
        net_pre = money(r[hdr["Net Price"]])
        cost = money(r[hdr["Actual Cost + 6% Tariff"]])
        rows.append(row(part, r[hdr["Description"]], "EA", money(r[hdr["MSRP"]]), cost if cost is not None else net_pre,
                        "", "", "", **{"Net Before Tariff": net_pre}))
    return rows, ["Net Before Tariff"], {"source": os.path.basename(path), "effective": "2026", "cost_basis": "Actual Cost + 6% Tariff"}


def parse_xcluder():
    """Xcluder: several 'Page N' sheets; section title rows, repeated 'Item #/Part # | Description | Xcluder Cost | Action Cost' headers."""
    path = find_raw("*XCLUDER*.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows, cat = [], ""
    m = re.search(r"\((\d{2})-(\d{2})-(\d{4})\)", os.path.basename(path))
    eff = f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else "2026"
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            cells = [c for c in r if c is not None and clean(c)]
            if not cells:
                continue
            first = clean(cells[0])
            if first.lower().startswith(("item #", "part #")):
                continue  # column header
            if len(cells) == 1:
                if first.startswith("**") or first.lower().startswith("kitted") or "price list" in first.lower():
                    continue
                cat = re.sub(r"[®™]", "", first)
                continue
            if len(cells) >= 3 and money(cells[2]) is not None:
                notes = ""
                if "MOQ" in cat or "min 20" in cat or "BULK" in cat.upper():
                    notes = "Bulk / minimum-order pricing tier"
                rows.append(row(cells[0], cells[1], "EA", None, money(cells[2]), cat, "", notes))
    return rows, [], {"source": os.path.basename(path), "effective": eff, "cost_basis": "Xcluder Cost"}


def parse_horton():
    """Horton via Advanced Door Services: PART | A&M Part # | Description | UOM | List Price | Our cost | | Notes."""
    path = find_raw("*HORTON*.xlsm") or find_raw("*HORTON*.xlsx")
    rows, hdr = [], None
    for r in sheet_rows(path):
        cells = [clean(c) for c in r]
        if hdr is None:
            if "PART" in cells and "Our cost" in cells:
                hdr = {name: cells.index(name) for name in cells if name}
            continue
        part = r[hdr["PART"]]
        if not clean(part):
            continue
        rows.append(row(part, r[hdr["Description"]], r[hdr["UOM"]], money(r[hdr["List Price"]]), money(r[hdr["Our cost"]]),
                        "", "", r[hdr.get("Notes", -1)] if "Notes" in hdr else "",
                        **{"A&M Part #": clean(r[hdr["Automatics & More Part #"]]) if "Automatics & More Part #" in hdr else ""}))
    m = re.search(r"\((\d{2})-(\d{2})-(\d{4})\)", os.path.basename(path))
    eff = f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else "2026"
    return rows, ["A&M Part #"], {"source": os.path.basename(path), "effective": eff, "cost_basis": "Our cost (30% off list)"}


def parse_record():
    """Record USA: 'Summary' sheet = Part | Description | UOM | List Price | Our Cost | Addison | DC | Best Cost."""
    path = find_raw("RECORD*.xlsx")
    rows, hdr = [], None
    for r in sheet_rows(path, "Summary"):
        cells = [clean(c) for c in r]
        if hdr is None:
            if "Part" in cells and "Our Cost" in cells:
                hdr = {name: cells.index(name) for name in cells if name}
            continue
        part = r[hdr["Part"]]
        if not clean(part):
            continue
        rows.append(row(part, r[hdr["Description"]], r[hdr["UOM"]], money(r[hdr["List Price"]]), money(r[hdr["Our Cost"]]),
                        "", "", "",
                        **{"Addison Price": money(r[hdr["Addison"]]), "Door Controls Price": money(r[hdr["DC"]]),
                           "Best Source": clean(r[hdr["Best Cost"]])}))
    return rows, ["Addison Price", "Door Controls Price", "Best Source"], \
        {"source": os.path.basename(path), "effective": "2025-04-01", "cost_basis": "Our Cost (35% off + 4)"}


def parse_ready_access():
    """Ready Access: Used On | Description | Part Number | Authorized Agent | Suggested Resale | List | (link).
    The same part repeats once per model it's used on — merge those into one row listing every model."""
    path = find_raw("*READY ACCESS*.xlsx")
    by_part, hdr = {}, None
    for r in sheet_rows(path, "Sheet1"):
        cells = [clean(c) for c in r]
        if hdr is None:
            if "Part Number" in cells and "Authorized Agent" in cells:
                hdr = {name: cells.index(name) for name in cells if name}
            continue
        part = part_key(r[hdr["Part Number"]])
        if not part:
            continue
        used_on = clean(r[hdr["Used On"]])
        if part in by_part:
            ex = by_part[part]
            if used_on and used_on not in ex["_used"]:
                ex["_used"].append(used_on)
            if ex["Net Price"] is None:
                ex["Net Price"] = money(r[hdr["Authorized Agent"]])
            continue
        desc = clean(r[hdr["Description"]])
        notes = "Discontinued / no stock" if re.search(r"NO MORE", desc, re.I) else ""
        link = next((clean(c) for c in r[len(hdr):] if c and str(c).startswith("http")), "")
        d = row(part, desc, "EA", money(r[hdr["List"]]), money(r[hdr["Authorized Agent"]]), "", "", notes,
                **{"Suggested Resale": money(r[hdr["Suggested Resale"]]), "Link": link})
        d["_used"] = [used_on] if used_on else []
        by_part[part] = d
    rows = []
    for d in by_part.values():
        used = d.pop("_used")
        d["Category"] = ("Used on " + ", ".join(used)) if used else ""
        d["Used On"] = ", ".join(used)
        rows.append(d)
    return rows, ["Used On", "Suggested Resale", "Link"], {"source": os.path.basename(path), "effective": "2026", "cost_basis": "Authorized Agent"}


def parse_optex():
    """OPTEX level 5: Product Category | Division/Year | Part Number | Product Name | Description | LIST PRICE 2026 | NET (5).
    Category rows carry only col A; product-family rows carry a name but no part number."""
    path = find_raw("*OPTEX*.xlsx")
    rows, cat, family = [], "", ""
    for i, r in enumerate(sheet_rows(path)):
        if i == 0:
            continue
        c_cat, div, part, name, desc, lst, net = (list(r) + [None] * 7)[:7]
        if clean(c_cat):
            cat = clean(c_cat)
        if not clean(part):
            if clean(name):
                family = clean(name)
            continue
        full = clean(name) + (" — " + clean(desc) if clean(desc) else "")
        rows.append(row(part, full, "EA", money(lst), money(net), cat, "", "",
                        **{"Model": clean(name), "Product Family": family, "Division": clean(div)}))
    return rows, ["Model", "Product Family", "Division"], {"source": os.path.basename(path), "effective": "2026-01-02", "cost_basis": "NET (level 5)"}


def parse_ms_sedco():
    """MS Sedco: a PDF-to-Excel conversion. Header rows ('PART #', 'DESCRIPTION', 'RETAIL', 'YOUR PRICE') repeat with
    shifting column positions; single-cell rows between them are product-family headings."""
    path = find_raw("*MS SEDCO*.xlsx")
    rows, cols, cat = [], None, ""
    boiler = re.compile(r"click on images|order today|call 317|mssedco\.com|confidential|contents|how to order", re.I)
    for r in sheet_rows(path):
        cells = [clean(c) for c in r]
        if "PART #" in cells:
            cols = {n: cells.index(n) for n in ("PART #", "DESCRIPTION", "RETAIL", "YOUR PRICE") if n in cells}
            continue
        nonempty = [(i, c) for i, c in enumerate(cells) if c]
        if not nonempty or cols is None:
            continue
        if len(nonempty) == 1:
            txt = nonempty[0][1]
            raw = str(r[nonempty[0][0]])
            looks_like_heading = (not boiler.search(txt) and len(txt) < 70 and "\n" not in raw
                                  and not txt.endswith((".", ":")) and not re.fullmatch(r"[\d\s\-]+", txt)
                                  and not all(re.fullmatch(r"[A-Z0-9\-/]+", t) for t in txt.split()))  # part-number captions
            if looks_like_heading:
                cat = re.sub(r"[®™]", "", txt)
            continue
        part = cells[cols["PART #"]] if cols["PART #"] < len(cells) else ""
        if not part:
            continue
        desc = cells[cols["DESCRIPTION"]] if "DESCRIPTION" in cols else ""
        lst = money(cells[cols["RETAIL"]]) if "RETAIL" in cols else None
        net = money(cells[cols["YOUR PRICE"]]) if "YOUR PRICE" in cols else None
        if lst is None and net is None:
            continue
        rows.append(row(part, desc, "EA", lst, net, cat, "", ""))
    return rows, [], {"source": os.path.basename(path), "effective": "2026", "cost_basis": "YOUR PRICE"}


def parse_nabco():
    """NABCO 'Parts Look Up' PDF: NEW Pt # | OLD Pt # | DESCRIPTION | UNIT | LIST PRICE | NET PRICE.
    The NEW part column is 16 chars wide on every page; the OLD part (when present) starts at col 16.
    Units are EA, AU, or inch marks (\" / \"2) for stock lengths/sheets."""
    path = find_raw("*NABCO PRICE LIST*.pdf")
    text = pdftotext(path, layout=True)
    rows, warn = [], []
    tail = re.compile(r"^(.*?)\s+(\S{1,4})\s+\$?([\d,]+\.\d\d)\s+\$?([\d,]+\.\d\d)\s*$")
    unit_map = {'"': "IN", '"2': "SQ IN"}
    for line in text.splitlines():
        if "NEW Pt #" in line or "Parts Look Up" in line or "$" not in line:
            continue
        new = line[:16].strip()
        if not new:
            continue  # stray $0.00 rows carry no part number
        m = tail.match(line[16:])
        if not m:
            warn.append(line.strip()[:80])
            continue
        body, uom, lst, net = m.groups()
        old = ""
        if body and body[0] != " ":
            parts = re.split(r"\s{2,}", body.strip(), maxsplit=1)
            old, body = (parts[0], parts[1]) if len(parts) == 2 else ("", parts[0])
        rows.append(row(new, body, unit_map.get(uom, uom), money(lst), money(net), "", "", "", **{"Old Part #": old}))
    meta = {"source": os.path.basename(path), "effective": "2026-05-19", "cost_basis": "NET PRICE",
            "warnings": [f"{len(warn)} priced lines did not parse: " + " | ".join(warn[:5])] if warn else []}
    try:
        decals = parse_nabco_decals()
    except Exception as e:  # OCR needs macOS (swift + Vision); on Windows/Linux just skip the decal pages
        decals = []
        meta["warnings"].append(f"decals PDF skipped — OCR unavailable on this machine ({type(e).__name__}); run on a Mac to include them")
    if decals:
        rows += decals
        meta["source"] += " + decals PDF (OCR)"
        meta["warnings"].append(f"{len(decals)} decal/module parts came from OCR of the image-only decals PDF — spot-check before quoting")
    return rows, ["Old Part #"], meta


# Lines the OCR can't read cleanly (an image overlaps the text) — transcribed by hand from the page.
OCR_FIXES = {
    "C-00118": 'Decal, "CAUTION STAND CLEAR", Two sided with wording on both sides. Overall size = 3.0" x 11.0"',
}


def parse_nabco_decals():
    """NABCO 'Modules and Decals' pages are an image-only PDF (broken font encoding), so we OCR them.
    Layout: ITEM picture | DESCRIPTION (x≈0.2–0.6) | PART # (x≈0.74) with the price to its right.
    Each description line and price is attached to the nearest part number vertically; parts stacked
    in one cell (several decals sharing one price) inherit the cell's price. Prices on these pages are LIST."""
    path = find_raw("*NABCO*DECALS*.pdf")
    if not path:
        return []
    words = ocr_pdf(path)
    part_rx = re.compile(r"^(?:[A-Z]-\d{5}|\d{2}-\d{4,5})$")
    price_rx = re.compile(r"^\$\s?[\d,]+\.\d\d$")
    rows = []
    embedded_rx = re.compile(r"\b([A-Z])[\-.](\d{5})\b")
    for page in sorted({w["page"] for w in words}):
        pw = [w for w in words if w["page"] == page]
        hdr = next((w for w in pw if w["text"].strip().upper().startswith("PART #")), None)
        px = hdr["x"] if hdr else 0.74  # part-number column moves between pages; anchor on its header
        dhdr = next((w for w in pw if w["text"].strip().upper() == "DESCRIPTION"), None)
        dx = (dhdr["x"] - 0.22) if dhdr else 0.15  # description text starts ~0.2 left of its centered header; pictures sit left of that
        parts = []
        for w in pw:
            t = w["text"].strip()
            if abs(w["x"] - px) < 0.08 and part_rx.match(t):
                parts.append(dict(w, text=t))
            elif w["x"] < px - 0.1:
                # OCR occasionally glues the part number onto the end of a description line ("... C.00118")
                m = embedded_rx.search(t)
                if m and t.upper() != t and len(t) > 20:
                    parts.append(dict(w, text=f"{m.group(1)}-{m.group(2)}", x=px))
                    w["text"] = t[:m.start()].rstrip(" :.-")
        parts.sort(key=lambda w: w["y"])
        if not parts:
            continue
        nearest = lambda y: min(parts, key=lambda p: abs(p["y"] - y))
        desc, price, note = {id(p): [] for p in parts}, {}, {id(p): [] for p in parts}
        for w in sorted(pw, key=lambda w: (w["y"], w["x"])):
            t = w["text"].strip()
            if not t:
                continue
            if dx < w["x"] < px - 0.16 and not part_rx.match(t) and t not in ("DESCRIPTION", "Section D1") \
                    and not t.startswith(("Modules and", "Decals")):
                desc[id(nearest(w["y"]))].append(t)
            elif px - 0.16 <= w["x"] < px - 0.02 and t == "Canadian":
                note[id(nearest(w["y"]))].append("Canadian version")
            elif w["x"] > px + 0.02 and price_rx.match(t):
                price[id(nearest(w["y"]))] = money(t)
        # stacked parts within one cell share the cell's price
        for i, p in enumerate(parts):
            if id(p) in price:
                continue
            j = i
            while j > 0 and parts[j]["y"] - parts[j - 1]["y"] < 0.025:
                j -= 1
            k = i
            while k < len(parts) - 1 and parts[k + 1]["y"] - parts[k]["y"] < 0.025:
                k += 1
            shared = [price[id(q)] for q in parts[j:k + 1] if id(q) in price]
            if shared:
                price[id(p)] = shared[0]
        for p in parts:
            d = OCR_FIXES.get(p["text"].strip()) or " ".join(desc[id(p)])
            if not d:
                continue
            # The decals pages print LIST price (M-01848 shows $25.00 there vs list $25 / net $12.50 on the
            # main list). NABCO's net is 50% of list across the main list, so derive our cost the same way.
            lst = price.get(id(p))
            net = round(lst * 0.5, 2) if lst is not None else None
            rows.append(row(p["text"].strip(), d, "EA", lst, net, "Modules & Decals", "",
                            "; ".join(note[id(p)] + ["OCR-derived; net = 50% of list (NABCO standard)"]), **{"Old Part #": ""}))
    return rows


def parse_motion_access():
    """Motion Access PDF (raw text order). Item lines begin with the 'Automatics & More Inc.' watermark and end with
    '<um> <price|CALL|-> <catalog page>'; long descriptions wrap onto following lines. Section rows: 'MotionAccess, LLC · X'."""
    path = find_raw("*MOTION ACCESS*.pdf")
    text = pdftotext(path, layout=False)
    rows, cat, buf = [], "", None
    tail = re.compile(r"^(.*?)\s+(each|feet|foot|pair|set|ft|kit|roll|box|lot)\s+(\$\s?[\d,]+(?:\.\d+)?|\$[\d,]+-[\d,]+|CALL|-)\s+(\d+)\s*$", re.I)
    part_rx = re.compile(r"^(option <[^>]+>|[A-Z0-9][A-Z0-9\-\./\*]*(?:\s\((?:single|pair)\*\))?|[A-Z0-9][A-Z0-9\-\./\*]*)\s+(.*)$")

    def flush(buf):
        m = tail.match(buf)
        if not m:
            return
        body, uom, price, page = m.groups()
        pm = part_rx.match(body)
        if not pm:
            return
        part, desc = pm.groups()
        notes = []
        net = None
        if price.upper() == "CALL":
            notes.append("Call for price")
        elif price != "-":
            rng = re.match(r"\$([\d,]+)-([\d,]+)$", price)
            if rng:
                net = money(rng.group(1))
                notes.append(f"Price range {price} depending on suffix")
            else:
                net = money(price)
        if re.search(r"core due", desc, re.I):
            notes.append("Core charge due")
        rows.append(row(part, desc, uom, None, net, cat, "", "; ".join(notes), **{"Catalog Page": page}))

    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("MotionAccess, LLC ·"):
            if buf:
                flush(buf); buf = None
            cat = s.split("·", 1)[1].strip()
            continue
        if s.startswith("Automatics & More Inc. "):
            if buf:
                flush(buf)
            buf = s[len("Automatics & More Inc. "):]
            if tail.match(buf):
                flush(buf); buf = None
            continue
        if buf is not None:
            buf = buf + " " + s
            if tail.match(buf):
                flush(buf); buf = None
    if buf:
        flush(buf)
    return rows, ["Catalog Page"], {"source": os.path.basename(path), "effective": "2026-05-26", "cost_basis": "Net Price",
                                    "warnings": ["Suffix-dependent items carry the LOW end of the vendor's price range; see Notes"]}


def parse_quad_systems():
    """Quad Systems sends a SALES QUOTE, not a price list: Item No. | Description | Unit | OEM | Qty | Unit Price | Total."""
    path = find_raw("*QUAD SYSTEMS*.pdf")
    text = pdftotext(path, layout=True)
    rows, cur = [], None
    item = re.compile(r"^\s?(\S+)\s{2,}(.+?)\s+(Each|EA|FT|Pair)\s+(\S*)\s+(\d+)\s+([\d,]+\.\d\d)\s+([\d,]+\.\d\d)\s*$")
    for line in text.splitlines():
        m = item.match(line)
        if m:
            part, desc, uom, oem, qty, unit, total = m.groups()
            if part.upper() == "TARIFF":
                cur = None
                continue
            cur = row(part, desc, uom, None, money(unit), oem, "", "From sales quote SQ608002 (01/20/26)", **{"OEM": oem})
            rows.append(cur)
            continue
        if cur and line.startswith(" " * 18) and line.strip() and not re.search(r"\d+\.\d\d\s*$", line) \
                and not re.match(r"\s*(Quad Systems|\d+ Navigation|Corpus Christi|\d{3}-\d{3}|Email:|Sell|To:|Ship|Terms|Item No)", line):
            cur["Description"] = clean(cur["Description"] + " " + line.strip())
        else:
            cur = None
    return rows, ["OEM"], {"source": os.path.basename(path), "effective": "2026-01-20", "cost_basis": "Quote unit price",
                           "warnings": ["Derived from a 14-line sales quote, not a full price list"]}


VENDORS = {
    "bea": parse_bea,
    "xcluder": parse_xcluder,
    "horton": parse_horton,
    "record-usa": parse_record,
    "ready-access": parse_ready_access,
    "optex": parse_optex,
    "ms-sedco": parse_ms_sedco,
    "nabco-gyrotech": parse_nabco,
    "motion-access": parse_motion_access,
    "quad-systems": parse_quad_systems,
}

SKIPPED = {}  # nothing skipped: the image-only NABCO decals PDF goes through macOS Vision OCR (see parse_nabco_decals)


def main(argv):
    wanted = [a for a in argv if a in VENDORS]
    unknown = [a for a in argv if a not in VENDORS]
    if unknown:
        print("unknown vendor id(s):", ", ".join(unknown), "\nknown:", ", ".join(VENDORS))
        return 2
    manifest_path = os.path.join(OUT_DIR, "manifest.json")
    manifest = {"generatedAt": date.today().isoformat(), "vendors": {}, "skipped": SKIPPED}
    if os.path.exists(manifest_path):
        try:
            manifest["vendors"] = json.load(open(manifest_path)).get("vendors", {})
        except Exception:
            pass
    for vid, fn in VENDORS.items():
        if wanted and vid not in wanted:
            continue
        try:
            rows, extras, meta = fn()
        except Exception as e:  # keep going; report at the end
            print(f"  ✗ {vid}: {e}")
            manifest["vendors"][vid] = {"error": str(e)}
            continue
        path, n, dupes = write_csv(vid, rows, extras)
        with_net = sum(1 for r in rows if r.get("Net Price") is not None)
        with_list = sum(1 for r in rows if r.get("List Price") is not None)
        manifest["vendors"][vid] = {**meta, "file": os.path.relpath(path, ROOT), "parts": n, "withNet": with_net,
                                    "withList": with_list, "duplicatesDropped": dupes}
        print(f"  ✓ {vid:15} {n:5} parts  net:{with_net:5}  list:{with_list:5}  dupes:{dupes:3}  ← {meta['source']}")
    os.makedirs(OUT_DIR, exist_ok=True)
    json.dump(manifest, open(manifest_path, "w"), indent=2)
    for k, v in SKIPPED.items():
        print(f"  – skipped {k}: {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
